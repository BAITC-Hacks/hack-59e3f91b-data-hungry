"""Generate the small sample uploads used by tests/test_attachments.py (idempotent).

Run:  uv run python tests/fixtures/make_fixtures.py
"""
from __future__ import annotations

import io
from pathlib import Path

FIXTURES = Path(__file__).resolve().parent

SPEC_ROWS = [
    ("№", "Артикул", "Наименование", "Кол-во", "Ед.", "Цена"),
    (1, "027228", "АВ DRX250 MT 3ф 160А 18ka Legrand", 2, "шт", 64920),
    (2, "ярп4520", "LED STARK 30W 2400Lm d98x180 4000K IP20 MEGALIGHT", 10, "шт", 1810),
    (3, "221-413", "Клемма соед. 3-пров. (0,2-4мм² 32А) WAGO", 50, "шт", 180),
    (4, None, "ВВГ п нг (А) 3 х 2,5 0,66 кВ ГОСТ EKT", 300, "м", 250),
    (5, "310100080_", "Диф.авт. 1p+N 16А (30мА) 411002 Legrand", 4, "шт", 9430),
    (6, "RMNF22TB30", "Реле контроля 3-фазное NFC 8A 2CO", 1, "шт", 45000),
]


def make_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Спецификация"
    ws.append(["Спецификация к запросу №17 от 23.09.2026"])
    ws.append([])
    for row in SPEC_ROWS:
        ws.append(list(row))
    ws2 = wb.create_sheet("Контакты")
    ws2.append(["Менеджер", "Иванов И.И."])
    wb.save(path)


def make_docx(path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_paragraph("Заявка на электротехническую продукцию")
    doc.add_paragraph("Просим выставить счёт на следующие позиции:")
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Артикул", "Наименование", "Кол-во, шт"
    for _n, art, name, qty, _unit, _price in SPEC_ROWS[1:4]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = art or "", name, str(qty)
    doc.add_paragraph("Доставка: г. Алматы, самовывоз.")
    doc.save(path)


def make_pdf(path: Path) -> None:
    """Minimal single-page PDF with a Latin-only text stream (pypdf can extract it)."""
    lines = [
        "SPECIFICATION 2026-09-23",
        "1 | 027228 | AV DRX250 MT 3f 160A Legrand | 2 pcs",
        "2 | 221-413 | WAGO 3-wire terminal | 50 pcs",
        "3 | LC3-C5E04-359 | ITK cable F/UTP cat.5E 305m | 1 pcs",
    ]
    content = "BT /F1 12 Tf 40 780 Td 16 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(out.getvalue())


def make_png(path: Path) -> None:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (2000, 1200), (245, 245, 240))
    d = ImageDraw.Draw(img)
    d.rectangle((200, 200, 1800, 1000), outline=(30, 30, 30), width=8)
    d.text((260, 260), "027228 Legrand DRX250", fill=(0, 0, 0))
    img.save(path, "PNG")


def make_all() -> dict[str, Path]:
    targets = {
        "xlsx": (FIXTURES / "spec.xlsx", make_xlsx),
        "docx": (FIXTURES / "spec.docx", make_docx),
        "pdf": (FIXTURES / "spec.pdf", make_pdf),
        "png": (FIXTURES / "photo.png", make_png),
    }
    for _k, (path, fn) in targets.items():
        if not path.exists():
            fn(path)
    return {k: p for k, (p, _f) in targets.items()}


if __name__ == "__main__":
    for k, p in make_all().items():
        print(k, p, p.stat().st_size, "bytes")
